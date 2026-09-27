# Categorize Media 對話框：依 YOLO/SAM3 偵測結果分類媒體檔案
# 輸出方式可選只產生 CSV / Excel / SQLite 索引檔 (不動原始檔案), 或搬移到子資料夾
# 更新日期: 2026-09-27
from __future__ import annotations

import csv
import shutil
import sqlite3
from collections import Counter
from pathlib import Path
from typing import ClassVar

import cv2
import orjson
from openpyxl import Workbook
from openpyxl.styles import Font
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QApplication,
    QComboBox,
    QDialog,
    QFileDialog,
    QFormLayout,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QVBoxLayout,
)

from src.utils.const import ALL_EXTS, IMAGE_EXTS, VIDEO_EXTS
from src.utils.dynamic_settings import settings
from src.utils.func import getXmlPath, imread_unicode
from src.utils.img_handler import sam3_label_conf
from src.utils.logger import getUniqueLogger

log = getUniqueLogger(__file__)

# 單筆偵測結果: (檔案路徑, 代表類別的原始 class name, {class_name: 偵測次數})
DetectResult = tuple[Path, str, dict[str, int]]


class CategorizeMediaDialog(QDialog):
    """依 YOLO 偵測結果分類媒體檔案：輸出 CSV / Excel / SQLite 索引, 或搬移到子資料夾"""

    DEFAULT_MODEL = "yolo26s.pt"
    NOT_DETECTED_FOLDER = "not_detected"
    FALLBACK_FOLDER = "unknown"
    VIDEO_SAMPLE_FRAMES = 5
    # Windows 檔名不允許的字元; SAM3 的類別名是使用者自由輸入的 text prompt
    INVALID_NAME_CHARS = '<>:"/\\|?*'
    # 索引檔固定產在目標資料夾內, 不另外跳存檔對話框
    RESULT_CSV_NAME = "categorize_result.csv"
    RESULT_XLSX_NAME = "categorize_result.xlsx"
    RESULT_DB_NAME = "categorize_result.db"
    RESULT_TABLE = "categorize_result"
    RESULT_FIELDS = (
        "file_name", "file_path", "category",
        "detections", "total_count", "media_type",
    )
    # 輸出方式 → 索引檔檔名
    RESULT_NAMES: ClassVar[dict[str, str]] = {
        "csv": RESULT_CSV_NAME,
        "excel": RESULT_XLSX_NAME,
        "sqlite": RESULT_DB_NAME,
    }

    def __init__(
        self, parent=None, default_folder: str = "", default_model: str = ""
    ):
        super().__init__(parent)
        self.setWindowTitle("Categorize Media")
        self.setMinimumWidth(500)
        self._canceled = False
        # 最近一次判斷過類型的 model 路徑, 用來省下重複的 torch.load
        self._detected_model_path = ""

        main_layout = QVBoxLayout(self)

        # 說明
        hint = QLabel(
            "使用 YOLO 模型偵測資料夾中的圖片與影片，\n"
            "依偵測到最多次的物件名稱決定每個檔案的分類\n"
            "（也可使用 SAM3 model，但分類效果通常不如 YOLO）"
        )
        hint.setStyleSheet("color: gray; font-size: 11px;")
        hint.setWordWrap(True)
        main_layout.addWidget(hint)

        # --- 資料夾選擇 ---
        form = QFormLayout()
        folder_row = QHBoxLayout()
        # 路徑可直接打字或貼上, 不一定要走「瀏覽...」
        self.folder_edit = QLineEdit(default_folder)
        self.folder_edit.setPlaceholderText("選擇或直接輸入要分類的資料夾路徑")
        self.folder_edit.setToolTip("可直接輸入或貼上路徑，也可按「瀏覽...」選擇")
        self.folder_edit.textChanged.connect(self._update_output_hint)
        self.folder_edit.editingFinished.connect(self._on_folder_edited)
        folder_browse = QPushButton("瀏覽...")
        folder_browse.setFixedWidth(80)
        folder_browse.clicked.connect(self._browse_folder)
        folder_row.addWidget(self.folder_edit)
        folder_row.addWidget(folder_browse)
        form.addRow("資料夾:", folder_row)

        # --- Model 選擇 ---
        model_row = QHBoxLayout()
        self.type_combo = QComboBox()
        self.type_combo.addItem("YOLO", "yolo")
        self.type_combo.addItem("YOLO-Seg", "yolo-seg")
        self.type_combo.addItem("SAM3", "sam3")
        self.type_combo.setFixedWidth(100)
        # model 路徑同樣可直接打字或貼上
        self.model_edit = QLineEdit(default_model)
        self.model_edit.setPlaceholderText("選擇或直接輸入用於分類的 model (.pt)")
        self.model_edit.setToolTip(
            "可直接輸入或貼上 .pt 路徑，也可按「瀏覽...」選擇\n"
            "輸入完成後會自動判斷模型類型 (YOLO / YOLO-Seg / SAM3)"
        )
        self.model_edit.editingFinished.connect(self._on_model_edited)
        model_browse = QPushButton("瀏覽...")
        model_browse.setFixedWidth(80)
        model_browse.clicked.connect(self._browse_model)
        model_reset = QPushButton("Reset")
        model_reset.setFixedWidth(60)
        model_reset.setToolTip(f"重設為預設模型 ({self.DEFAULT_MODEL})")
        model_reset.clicked.connect(self._reset_model)
        model_row.addWidget(self.type_combo)
        model_row.addWidget(self.model_edit)
        model_row.addWidget(model_browse)
        model_row.addWidget(model_reset)
        form.addRow("Model:", model_row)

        # --- 輸出方式 ---
        self.output_combo = QComboBox()
        # 搬移是不可逆的, 排在最後一項; 預設落在不動原始檔案的 CSV
        self.output_combo.addItem(f"產生 CSV 檔 ({self.RESULT_CSV_NAME})", "csv")
        self.output_combo.addItem(
            f"產生 Excel 檔 ({self.RESULT_XLSX_NAME})", "excel"
        )
        self.output_combo.addItem(
            f"產生 SQLite 檔 ({self.RESULT_DB_NAME})", "sqlite"
        )
        self.output_combo.addItem("搬移到子資料夾 (不可逆)", "move")
        output_tips = (
            "只產生 CSV 索引檔, 原始檔案留在原地",
            "只產生 Excel 索引檔 (內容同 CSV), 原始檔案留在原地",
            "只產生 SQLite 索引檔, 原始檔案留在原地",
            "把檔案搬到以類別命名的子資料夾, 原始檔案位置會改變且無法還原",
        )
        for i, tip in enumerate(output_tips):
            self.output_combo.setItemData(i, tip, Qt.ItemDataRole.ToolTipRole)
        self.output_combo.setToolTip(
            "索引檔記錄每個檔案的分類與各類別偵測次數, 不搬動原始檔案"
        )
        self.output_combo.currentIndexChanged.connect(self._update_output_hint)
        form.addRow("輸出方式:", self.output_combo)
        main_layout.addLayout(form)

        # 依「資料夾 + 輸出方式」顯示實際產出位置, 按下去之前就看得到
        self.output_hint = QLabel()
        self.output_hint.setStyleSheet("color: gray; font-size: 11px;")
        self.output_hint.setWordWrap(True)
        main_layout.addWidget(self.output_hint)
        self._update_output_hint()

        # --- 按鈕 ---
        btn_layout = QHBoxLayout()
        btn_layout.addStretch()
        self.start_btn = QPushButton("開始偵測")
        self.start_btn.clicked.connect(self._run)
        self.close_dialog_btn = QPushButton("關閉")
        self.close_dialog_btn.clicked.connect(self._on_cancel)
        btn_layout.addWidget(self.start_btn)
        btn_layout.addWidget(self.close_dialog_btn)
        main_layout.addLayout(btn_layout)

        # --- 進度區域（類似狀態列）---
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        self.progress_bar.setTextVisible(True)
        main_layout.addWidget(self.progress_bar)

        self.status_label = QLabel("")
        self.status_label.setStyleSheet("color: gray; font-size: 11px;")
        main_layout.addWidget(self.status_label)

    def _browse_folder(self):
        path = QFileDialog.getExistingDirectory(
            self, "選擇資料夾", self._clean_path_text(self.folder_edit.text())
        )
        if path:
            # setText 會觸發 textChanged, 提示列自己會更新
            self.folder_edit.setText(path)

    def _browse_model(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "選擇 Model",
            self._clean_path_text(self.model_edit.text()),
            "Model Files (*.pt)",
        )
        if path:
            self.model_edit.setText(path)
            self._apply_model_type(path)

    def _on_model_edited(self):
        """model 路徑輸入完成後正規化, 並自動判斷模型類型"""
        cleaned = self._clean_path_text(self.model_edit.text())
        if cleaned != self.model_edit.text():
            self.model_edit.setText(cleaned)

        if not cleaned:
            self.status_label.setText("")
            return
        if not Path(cleaned).is_file():
            self.status_label.setText(f"⚠ 找不到這個 model 檔案：{cleaned}")
            return

        self.status_label.setText("")
        # torch.load 讀整個 checkpoint 並不便宜, 同一個路徑只判斷一次
        if cleaned != self._detected_model_path:
            self._apply_model_type(cleaned)

    def _apply_model_type(self, model_path: str):
        """判斷 model 類型並同步左側的類型下拉選單"""
        self._detected_model_path = model_path
        # 大模型載入要數秒, 先讓使用者知道畫面不是卡住
        self.status_label.setText("正在判斷模型類型...")
        QApplication.processEvents()
        detected = self._detect_model_type(model_path)
        self.status_label.setText("")
        if detected:
            idx = self.type_combo.findData(detected)
            if idx >= 0:
                self.type_combo.setCurrentIndex(idx)

    def _reset_model(self):
        """重設為預設 YOLO model"""
        self.model_edit.setText(self.DEFAULT_MODEL)
        self.type_combo.setCurrentIndex(0)  # YOLO
        # 預設就是 YOLO, 類型已經對了, 不必再 torch.load 判斷一次
        self._detected_model_path = self.DEFAULT_MODEL
        self.status_label.setText("")

    @staticmethod
    def _clean_path_text(text: str) -> str:
        """去掉路徑前後的空白與引號 (檔案總管的「複製路徑」會帶雙引號)"""
        return text.strip().strip('"').strip("'").strip()

    def _on_folder_edited(self):
        """輸入完成後把欄位內容正規化, 貼進來的引號不要留在畫面上"""
        cleaned = self._clean_path_text(self.folder_edit.text())
        if cleaned != self.folder_edit.text():
            self.folder_edit.setText(cleaned)  # 觸發 textChanged 更新提示列

    def _update_output_hint(self):
        """更新輸出位置提示文字"""
        folder = self._clean_path_text(self.folder_edit.text())
        mode = self.output_combo.currentData()
        # 路徑可手動輸入, 打錯字當場就要看得出來, 不必等按下「開始偵測」
        if folder and not Path(folder).is_dir():
            self.output_hint.setText(f"⚠ 找不到這個資料夾：{folder}")
            return

        folder = folder or "<資料夾>"
        # 用 Path 組合, 提示文字的分隔符才不會混用正反斜線
        if mode == "move":
            dest = Path(folder, "<類別>")
            self.output_hint.setText(
                f"檔案會搬移到 {dest} （原始檔案位置會改變, 無法還原）"
            )
        else:
            name = self.RESULT_NAMES[mode]
            self.output_hint.setText(
                f"索引檔：{Path(folder, name)} （原始檔案不會搬動）"
            )

    @staticmethod
    def _detect_model_type(model_path: str) -> str | None:
        """偵測 .pt 模型的類型 (yolo / yolo-seg / sam3)"""
        try:
            import torch
            ckpt = torch.load(model_path, map_location="cpu", weights_only=False)
            if isinstance(ckpt, dict) and "model" in ckpt:
                cls_name = type(ckpt["model"]).__name__.lower()
                if "sam" in cls_name:
                    return "sam3"
                # ultralytics 的 checkpoint 其實沒有 model.task (getattr 一律拿到
                # 空字串, seg model 因此被判成 yolo); 類別名稱 SegmentationModel
                # 才是直接的依據, task 則記在 train_args 裡
                if "segment" in cls_name:
                    return "yolo-seg"
                task = getattr(ckpt["model"], "task", "") or ""
                if not task:
                    task = (ckpt.get("train_args") or {}).get("task", "") or ""
                if task == "segment":
                    return "yolo-seg"
                return "yolo"
        except Exception:
            log.w(f"無法偵測模型類型: {model_path}")
        return None

    def _on_cancel(self):
        """取消按鈕：偵測中則中斷，否則關閉"""
        self._canceled = True
        self.reject()

    def _run(self):
        """開始偵測, 再依輸出方式搬移檔案或產生索引檔"""
        folder = self._clean_path_text(self.folder_edit.text())
        model_path = self._clean_path_text(self.model_edit.text())
        output_mode = self.output_combo.currentData()

        if not folder or not Path(folder).is_dir():
            QMessageBox.warning(self, "Warning", "請選擇或輸入有效的資料夾路徑")
            return
        if not model_path or not Path(model_path).is_file():
            QMessageBox.warning(
                self, "Warning", "請選擇或輸入有效的 Model 檔案 (.pt)"
            )
            return

        # 收集媒體檔案（不含子資料夾）
        base = Path(folder)
        media_files = sorted(
            f for f in base.iterdir()
            if f.is_file() and f.suffix.lower() in ALL_EXTS
        )
        if not media_files:
            QMessageBox.warning(self, "Warning", "資料夾中沒有找到圖片或影片檔案")
            return

        # 索引檔的覆蓋確認提前到偵測之前, 免得整輪跑完才發現使用者不想覆蓋
        out_path: Path | None = None
        if output_mode != "move":
            out_path = self._prepare_output_path(base, output_mode)
            if out_path is None:
                return

        # 載入 model
        model_type = self.type_combo.currentData()
        self.start_btn.setEnabled(False)
        self._canceled = False
        self.status_label.setText("正在載入模型...")
        QApplication.processEvents()

        sam3_labels: list[str] = []
        try:
            if model_type == "sam3":
                from ultralytics.models.sam import SAM3SemanticPredictor

                sam3_labels = list(
                    dict.fromkeys(settings.class_names.text_prompts or [])
                )
                if not sam3_labels:
                    QMessageBox.warning(
                        self, "Warning",
                        "SAM3 需要 Text Prompts 才能偵測，\n"
                        "請先在 Edit → Text Prompts 中設定",
                    )
                    self.start_btn.setEnabled(True)
                    return
                overrides = dict(
                    conf=settings.models.sam3_conf or 0.25,
                    imgsz=630, task="segment",
                    # quantize=16 即 FP16, 取代已 deprecated 的 half=True
                    mode="predict", model=model_path, quantize=16, verbose=False,
                )
                model = SAM3SemanticPredictor(overrides=overrides)
            else:
                from ultralytics import YOLO
                model = YOLO(model_path)
        except Exception:
            log.e(f"無法載入模型: {model_path}")
            QMessageBox.critical(self, "Error", "模型載入失敗，請確認檔案是否正確")
            self.start_btn.setEnabled(True)
            return

        # 偵測每個檔案
        total = len(media_files)
        self.progress_bar.setVisible(True)
        self.progress_bar.setMaximum(total)
        self.progress_bar.setValue(0)

        results: list[DetectResult] = []

        for i, file_path in enumerate(media_files):
            if self._canceled:
                break

            self.status_label.setText(
                f"偵測中: {file_path.name} ({i + 1}/{total})"
            )
            self.progress_bar.setValue(i)
            QApplication.processEvents()

            try:
                if model_type == "sam3":
                    class_counts, class_confs = self._detect_file_sam3(
                        model, file_path, sam3_labels
                    )
                else:
                    class_counts, class_confs = self._detect_file(model, file_path)
            except Exception:
                log.e(f"偵測失敗: {file_path.name}")
                class_counts, class_confs = {}, {}

            if not class_counts:
                category = self.NOT_DETECTED_FOLDER
            else:
                # 這裡留原始的 class name, 檔名淨化等到真的要 mkdir 時才做:
                # 索引檔是純文字欄位, 沒有檔案系統的限制, 記淨化過的名字反而失真
                category = self._top_class(class_counts, class_confs)

            results.append((file_path, category, class_counts))

        if self._canceled:
            self.status_label.setText("已取消")
            self.progress_bar.setVisible(False)
            self.start_btn.setEnabled(True)
            return

        # 依輸出方式處理結果
        if output_mode == "move":
            ok = self._output_move(base, results)
        else:
            ok = self._output_index(out_path, output_mode, results)

        self.progress_bar.setValue(total)
        self.status_label.setText("完成" if ok else "失敗")
        self.start_btn.setEnabled(True)

    def _prepare_output_path(self, base: Path, output_mode: str) -> Path | None:
        """回傳索引檔路徑, 已存在則先問是否覆蓋; 使用者取消回傳 None"""
        name = self.RESULT_NAMES[output_mode]
        out_path = base / name
        if not out_path.exists():
            return out_path

        if output_mode != "sqlite":
            msg = f"{name} 已存在，是否覆蓋？"
        else:
            msg = (
                f"{name} 已存在，是否覆蓋其中的 {self.RESULT_TABLE} 資料表？\n"
                "（同一個檔案內的其他資料表不受影響）"
            )
        reply = QMessageBox.question(
            self, "檔案已存在", msg,
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if reply != QMessageBox.StandardButton.Yes:
            return None
        return out_path

    def _output_move(self, base: Path, results: list[DetectResult]) -> bool:
        """把檔案搬到以類別命名的子資料夾, 顯示摘要並回傳是否全部成功"""
        self.status_label.setText("正在搬移檔案...")
        QApplication.processEvents()

        moved_counts: dict[str, int] = {}
        failed: list[str] = []
        for file_path, category, _counts in results:
            # 類別名到了檔案系統才需要淨化
            subfolder = self._folder_name(category)
            dest_dir = base / subfolder
            # 先算好標註路徑: 圖片搬走後就找不到原本的位置了
            xml_path = getXmlPath(file_path)
            try:
                dest_dir.mkdir(exist_ok=True)
                shutil.move(str(file_path), str(dest_dir / file_path.name))
            except Exception as e:
                # 單一檔案失敗不該中斷整批: 未捕捉的例外會被全域 excepthook 記下來後
                # 直接結束這個迴圈, 只留下已建好的空資料夾, 其餘檔案完全沒被搬移
                log.e(f"搬移失敗 ({file_path.name} -> {subfolder}/): {e}")
                failed.append(file_path.name)
                continue
            moved_counts[subfolder] = moved_counts.get(subfolder, 0) + 1

            # VOC 標註跟著圖片走; 留在原地的話圖片被分類後標註就斷開了。
            # 圖片已經搬成功, 標註搬失敗只記錄下來, 不影響這個檔案的分類結果
            if xml_path.is_file():
                try:
                    shutil.move(str(xml_path), str(dest_dir / xml_path.name))
                except Exception as e:
                    log.e(f"標註搬移失敗 ({xml_path.name} -> {subfolder}/): {e}")
                    failed.append(xml_path.name)

        lines = ["分類完成\n"]
        for subfolder in sorted(moved_counts.keys()):
            lines.append(f"  {subfolder}/: {moved_counts[subfolder]} 個檔案")
        lines.append(f"\n共處理 {len(results)} 個檔案")
        if failed:
            lines.append(f"搬移失敗 {len(failed)} 個檔案（詳見 log）")
        QMessageBox.information(self, "Categorize Media 結果", "\n".join(lines))
        return not failed

    def _output_index(
        self, out_path: Path, output_mode: str, results: list[DetectResult]
    ) -> bool:
        """產生 CSV / Excel / SQLite 索引檔 (不搬動原始檔案), 顯示摘要並回傳是否成功"""
        self.status_label.setText("正在寫入索引檔...")
        QApplication.processEvents()

        rows = self._build_rows(results)
        if output_mode == "csv":
            ok = self._write_csv(out_path, rows)
        elif output_mode == "excel":
            ok = self._write_excel(out_path, rows)
        else:
            ok = self._write_sqlite(out_path, rows)

        if not ok:
            QMessageBox.critical(
                self, "Error", f"索引檔寫入失敗：{out_path.name}\n詳細原因請見 log"
            )
            return False

        category_counts: dict[str, int] = {}
        for _file_path, category, _counts in results:
            category_counts[category] = category_counts.get(category, 0) + 1

        lines = [f"索引檔已產生：\n{out_path}\n"]
        for category in sorted(category_counts.keys()):
            lines.append(f"  {category}: {category_counts[category]} 個檔案")
        lines.append(f"\n共處理 {len(results)} 個檔案，原始檔案未搬動")
        QMessageBox.information(self, "Categorize Media 結果", "\n".join(lines))
        return True

    @staticmethod
    def _build_rows(results: list[DetectResult]) -> list[dict]:
        """把偵測結果整理成索引檔的資料列"""
        rows: list[dict] = []
        for file_path, category, counts in results:
            media_type = (
                "image" if file_path.suffix.lower() in IMAGE_EXTS else "video"
            )
            rows.append({
                "file_name": file_path.name,
                "file_path": str(file_path),
                "category": category,
                # 完整的 {class_name: 次數}, 保留 category 以外被偵測到的類別;
                # 事後想換條件重新篩選就不必再跑一次模型
                "detections": orjson.dumps(
                    counts, option=orjson.OPT_SORT_KEYS
                ).decode(),
                "total_count": sum(counts.values()),
                "media_type": media_type,
            })
        return rows

    def _write_csv(self, out_path: Path, rows: list[dict]) -> bool:
        """寫出 CSV 索引檔, 回傳是否成功"""
        try:
            # utf-8-sig: 帶 BOM, Excel 直接開中文類別名稱才不會亂碼
            with out_path.open("w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=list(self.RESULT_FIELDS))
                writer.writeheader()
                writer.writerows(rows)
        except Exception:
            log.e(f"CSV 索引檔寫入失敗: {out_path}")
            return False
        return True

    def _write_excel(self, out_path: Path, rows: list[dict]) -> bool:
        """寫出 Excel 索引檔 (欄位同 CSV), 回傳是否成功"""
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = self.RESULT_TABLE
            ws.append(list(self.RESULT_FIELDS))
            for row in rows:
                ws.append([row[k] for k in self.RESULT_FIELDS])
            # 標題列粗體並凍結, 加上篩選鈕, 開檔就能直接依 category 篩選
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            # 欄寬依內容最長者估算, 上限 80 以免 file_path 撐得太寬
            for col in ws.columns:
                width = max(len(str(c.value)) for c in col if c.value is not None)
                ws.column_dimensions[col[0].column_letter].width = min(width + 2, 80)
            # 檔案被 Excel 開著時這裡會 PermissionError, 交給下方 except 記 log
            wb.save(out_path)
        except Exception:
            log.e(f"Excel 索引檔寫入失敗: {out_path}")
            return False
        return True

    def _write_sqlite(self, out_path: Path, rows: list[dict]) -> bool:
        """寫出 SQLite 索引檔, 回傳是否成功"""
        conn = None
        try:
            conn = sqlite3.connect(str(out_path))
            cur = conn.cursor()
            # 重跑時整個 table 重建, 避免殘留上一次的資料列
            cur.execute(f"DROP TABLE IF EXISTS {self.RESULT_TABLE}")
            cur.execute(
                f"CREATE TABLE {self.RESULT_TABLE} ("
                "file_name TEXT NOT NULL, "
                "file_path TEXT NOT NULL, "
                "category TEXT NOT NULL, "
                "detections TEXT NOT NULL, "
                "total_count INTEGER NOT NULL, "
                "media_type TEXT NOT NULL)"
            )
            cur.execute(
                f"CREATE INDEX idx_{self.RESULT_TABLE}_category "
                f"ON {self.RESULT_TABLE}(category)"
            )
            placeholders = ", ".join(["?"] * len(self.RESULT_FIELDS))
            cur.executemany(
                f"INSERT INTO {self.RESULT_TABLE} "
                f"({', '.join(self.RESULT_FIELDS)}) VALUES ({placeholders})",
                [tuple(row[k] for k in self.RESULT_FIELDS) for row in rows],
            )
            conn.commit()
        except Exception:
            log.e(f"SQLite 索引檔寫入失敗: {out_path}")
            return False
        finally:
            if conn is not None:
                try:
                    conn.close()
                except Exception:
                    log.e(f"SQLite 連線關閉失敗: {out_path}")
        return True

    def _detect_file(
        self, model, file_path: Path
    ) -> tuple[dict[str, int], dict[str, float]]:
        """偵測單一檔案，回傳 ({class_name: 次數}, {class_name: 信心值總和})"""
        counts: Counter = Counter()
        confs: Counter = Counter()
        suffix = file_path.suffix.lower()

        if suffix in IMAGE_EXTS:
            img = imread_unicode(file_path)
            if img is not None:
                self._count_detections(model, img, counts, confs)
        elif suffix in VIDEO_EXTS:
            cap = cv2.VideoCapture(str(file_path))
            try:
                total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
                if total_frames > 0:
                    for idx in self._sample_frame_indices(total_frames):
                        cap.set(cv2.CAP_PROP_POS_FRAMES, idx)
                        ret, frame = cap.read()
                        if ret:
                            self._count_detections(model, frame, counts, confs)
            finally:
                cap.release()

        return dict(counts), dict(confs)

    @staticmethod
    def _count_detections(model, img, counts: Counter, confs: Counter):
        """對單一影像跑 YOLO 推論並累加 class_name 的次數與信心值"""
        conf = settings.models.yolo_conf or 0.25
        results = model.predict(img, conf=conf, verbose=False)
        for r in results:
            if r.boxes is not None:
                for box in r.boxes:
                    name = model.names[int(box.cls)]
                    counts[name] += 1
                    confs[name] += float(box.conf)

    def _detect_file_sam3(
        self, predictor, file_path: Path, labels: list[str]
    ) -> tuple[dict[str, int], dict[str, float]]:
        """SAM3 偵測單一檔案，回傳 ({class_name: 次數}, {class_name: 信心值總和})"""
        counts: Counter = Counter()
        confs: Counter = Counter()
        suffix = file_path.suffix.lower()

        if suffix in IMAGE_EXTS:
            img = imread_unicode(file_path)
            if img is not None:
                self._count_sam3(predictor, img, labels, counts, confs)
        elif suffix in VIDEO_EXTS:
            cap = cv2.VideoCapture(str(file_path))
            try:
                total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
                if total_frames > 0:
                    for idx in self._sample_frame_indices(total_frames):
                        cap.set(cv2.CAP_PROP_POS_FRAMES, idx)
                        ret, frame = cap.read()
                        if ret:
                            self._count_sam3(
                                predictor, frame, labels, counts, confs
                            )
            finally:
                cap.release()

        return dict(counts), dict(confs)

    @staticmethod
    def _count_sam3(
        predictor, img, labels: list[str], counts: Counter, confs: Counter
    ):
        """對單一影像跑 SAM3 推論並累加 class_name 的次數與信心值"""
        predictor.set_image(img)
        src_shape = img.shape[:2]
        masks, boxes = predictor.inference_features(
            predictor.features, src_shape=src_shape, text=labels
        )
        # boxes 為 (N, 6) = xyxy + score + cls, cls 是 text prompt 的索引 (非偵測序號)
        boxes_np = boxes.cpu().numpy() if boxes is not None else None
        if boxes_np is not None:
            for i, box in enumerate(boxes_np):
                x1, y1, x2, y2 = int(box[0]), int(box[1]), int(box[2]), int(box[3])
                if (x2 - x1) > 0 and (y2 - y1) > 0:
                    label, score = sam3_label_conf(boxes_np, i, labels)
                    counts[label] += 1
                    # 取不到分數時 sam3_label_conf 回傳 -1.0, 夾成 0 以免拉低總和
                    confs[label] += max(score, 0.0)

    @staticmethod
    def _top_class(counts: dict[str, int], confs: dict[str, float]) -> str:
        """挑出代表整個檔案的單一類別。

        一個檔案只進一個 class 資料夾: 先比偵測次數, 同票比信心值總和, 再同票
        取字母序較前者, 讓同一批檔案每次跑的結果一致。

        原本同票是把類別名用 `+` 串成資料夾名 (例如 `dog+person`), 但「各出現
        一次」的組合非常常見, 每種組合都會長出一個新資料夾, 結果幾乎是一張圖
        一個資料夾, 失去分類的意義。
        """
        return min(
            counts,
            key=lambda name: (-counts[name], -confs.get(name, 0.0), name),
        )

    @classmethod
    def _folder_name(cls, class_name: str) -> str:
        """把類別名轉成可用的資料夾名。

        SAM3 的類別來自使用者自由輸入的 text prompt, 可能含有 `/`、`?` 這類
        字元, 直接拿來 mkdir 會失敗或意外建出多層資料夾。
        """
        name = "".join(
            "_" if ch in cls.INVALID_NAME_CHARS or ord(ch) < 32 else ch
            for ch in class_name
        )
        # Windows 會忽略結尾的空白與點, 建出來的資料夾名會和預期不同
        name = name.strip().rstrip(". ")
        return name or cls.FALLBACK_FOLDER

    def _sample_frame_indices(self, total: int) -> list[int]:
        """從影片中均勻取樣 frame indices"""
        n = min(self.VIDEO_SAMPLE_FRAMES, total)
        if n <= 1:
            return [0]
        return [int(i * (total - 1) / (n - 1)) for i in range(n)]
